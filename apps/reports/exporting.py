"""Excel / CSV export of the price comparison report."""

from __future__ import annotations

import csv
import io
from decimal import Decimal

from django.utils import timezone
from django.utils.translation import gettext as _
from openpyxl import Workbook
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from apps.catalog.models import Product
from apps.pricing.services import recommend_price


def report_rows(org) -> list[dict]:
    rows = []
    products = (
        Product.objects.for_org(org).filter(is_active=True).select_related("category").prefetch_related("offers__store")
    )
    for p in products:
        offers = [o for o in p.offers.all() if o.is_active]
        market_min = p.market_min()
        in_stock_prices = [o.last_price for o in offers if o.last_price is not None and o.last_in_stock is not False]
        avg = sum(in_stock_prices, Decimal(0)) / len(in_stock_prices) if in_stock_prices else None
        rec = recommend_price(p)
        rows.append(
            {
                "sku": p.sku,
                "name": p.name,
                "category": p.category.name if p.category else "",
                "my_price": p.my_price,
                "cost_price": p.cost_price,
                "market_min": market_min,
                "market_avg": avg.quantize(Decimal("1")) if avg is not None else None,
                "diff_pct": ((p.my_price - market_min) / market_min * 100).quantize(Decimal("0.1"))
                if market_min and p.my_price
                else None,
                "recommended": rec.price,
                "competitors": {o.store.name: (o.last_price, o.last_in_stock) for o in offers},
            }
        )
    return rows


def _headers(stores: list[str]) -> list[str]:
    return [
        "SKU",
        _("Product"),
        _("Category"),
        _("My price"),
        _("Cost price"),
        _("Market min"),
        _("Market avg"),
        _("Diff vs min, %"),
        _("Recommended"),
    ] + stores


def _flat(row: dict, stores: list[str]) -> list:
    values = [
        row["sku"],
        row["name"],
        row["category"],
        row["my_price"],
        row["cost_price"],
        row["market_min"],
        row["market_avg"],
        row["diff_pct"],
        row["recommended"],
    ]
    for s in stores:
        price, in_stock = row["competitors"].get(s, (None, None))
        values.append(price if in_stock is not False else (f"{price} (out)" if price is not None else "out"))
    return [float(v) if isinstance(v, Decimal) else v for v in values]


def export_csv(org) -> bytes:
    rows = report_rows(org)
    stores = sorted({s for r in rows for s in r["competitors"]})
    buf = io.StringIO()
    writer = csv.writer(buf)
    writer.writerow(_headers(stores))
    for row in rows:
        writer.writerow(_flat(row, stores))
    return ("﻿" + buf.getvalue()).encode("utf-8")  # BOM so Excel opens UTF-8 correctly


def export_excel(org) -> bytes:
    rows = report_rows(org)
    stores = sorted({s for r in rows for s in r["competitors"]})
    wb = Workbook()
    ws = wb.active
    ws.title = _("Prices")
    headers = _headers(stores)
    ws.append(headers)
    header_fill = PatternFill("solid", fgColor="1E293B")
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    for row in rows:
        ws.append(_flat(row, stores))

    money_cols = [4, 5, 6, 7, 9] + list(range(10, 10 + len(stores)))
    for col in money_cols:
        for cell in ws.iter_cols(min_col=col, max_col=col, min_row=2):
            for c in cell:
                c.number_format = "#,##0"
    widths = [14, 42, 18, 14, 14, 14, 14, 12, 14] + [16] * len(stores)
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = "C2"
    ws.auto_filter.ref = ws.dimensions

    last = len(rows) + 1
    if last > 1:
        red = PatternFill("solid", fgColor="FECACA")
        green = PatternFill("solid", fgColor="BBF7D0")
        # "More expensive than the market" highlight on the my-price column
        ws.conditional_formatting.add(f"D2:D{last}", FormulaRule(formula=["AND(ISNUMBER($F2),$D2>$F2)"], fill=red))
        ws.conditional_formatting.add(f"D2:D{last}", FormulaRule(formula=["AND(ISNUMBER($F2),$D2<=$F2)"], fill=green))
        ws.conditional_formatting.add(f"H2:H{last}", FormulaRule(formula=["AND(ISNUMBER($H2),$H2>0)"], fill=red))

    info = wb.create_sheet(_("About"))
    info.append([_("Organization"), org.name])
    info.append([_("Generated"), timezone.localtime().strftime("%Y-%m-%d %H:%M")])
    info.append([_("Products"), len(rows)])
    info.append([_("Legend"), _("Red = my price is above the cheapest competitor")])
    info.column_dimensions["A"].width = 18
    info.column_dimensions["B"].width = 60

    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()
