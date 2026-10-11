"""Bulk import of products and competitor links from CSV / Excel."""

from __future__ import annotations

import csv
import io
from dataclasses import dataclass, field
from decimal import Decimal
from typing import BinaryIO

from django.core.exceptions import ValidationError
from django.core.validators import URLValidator
from django.db import transaction
from django.utils.translation import gettext as _
from openpyxl import Workbook, load_workbook

from apps.billing.services import LimitExceeded, check_offer_limit
from apps.catalog.models import Category, CompetitorOffer, Product, Store, normalize_domain
from apps.scraping.price_parser import parse_price

TEMPLATE_HEADERS = ["sku", "name", "category", "my_price", "cost_price", "url_1", "url_2", "url_3"]
TEMPLATE_EXAMPLE = [
    "TV-55-QLED",
    'Samsung 55" QLED Q60D',
    "TV",
    "7990000",
    "6900000",
    "https://uzum.uz/ru/product/televizor-samsung-55-123456",
    "",
    "",
]


@dataclass
class ImportReport:
    created_products: int = 0
    updated_products: int = 0
    created_offers: int = 0
    errors: list[tuple[int, str]] = field(default_factory=list)

    @property
    def ok(self) -> bool:
        return not self.errors


def template_xlsx() -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "products"
    ws.append(TEMPLATE_HEADERS)
    ws.append(TEMPLATE_EXAMPLE)
    for col, width in zip("ABCDEFGH", [16, 40, 16, 14, 14, 50, 50, 50], strict=False):
        ws.column_dimensions[col].width = width
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()


def template_csv() -> bytes:
    buf = io.StringIO()
    writer = csv.writer(buf)
    writer.writerow(TEMPLATE_HEADERS)
    writer.writerow(TEMPLATE_EXAMPLE)
    return ("﻿" + buf.getvalue()).encode("utf-8")


def read_rows(file: BinaryIO, filename: str) -> list[dict]:
    if filename.lower().endswith((".xlsx", ".xlsm")):
        wb = load_workbook(file, read_only=True, data_only=True)
        ws = wb.active
        rows = list(ws.iter_rows(values_only=True))
        if not rows:
            return []
        headers = [str(h or "").strip().lower() for h in rows[0]]
        return [
            {headers[i]: ("" if v is None else str(v).strip()) for i, v in enumerate(r) if i < len(headers)}
            for r in rows[1:]
        ]
    raw = file.read()
    text = raw.decode("utf-8-sig") if isinstance(raw, bytes) else raw
    dialect = csv.Sniffer().sniff(text.splitlines()[0] if text else ",", delimiters=",;\t")
    reader = csv.DictReader(io.StringIO(text), dialect=dialect)
    return [{(k or "").strip().lower(): (v or "").strip() for k, v in row.items()} for row in reader]


def _store_for(url: str) -> Store:
    store = Store.for_url(url)
    if store is None:
        domain = normalize_domain(url)
        adapter = Store.Adapter.UZUM if domain == "uzum.uz" else Store.Adapter.JSONLD
        store = Store.objects.create(
            name=domain,
            domain=domain,
            adapter=adapter,
            requires_browser=adapter == Store.Adapter.UZUM,
            notes="Auto-created by import — verify adapter in admin",
        )
    return store


def import_rows(org, rows: list[dict], default_interval: int = 360) -> ImportReport:
    report = ImportReport()
    validate_url = URLValidator(schemes=["http", "https"])
    for line_no, row in enumerate(rows, start=2):
        if not any(row.values()):
            continue
        sku, name = row.get("sku", ""), row.get("name", "")
        if not sku or not name:
            report.errors.append((line_no, _("SKU and name are required")))
            continue
        my_price = parse_price(row.get("my_price")) if row.get("my_price") else Decimal(0)
        cost_price = parse_price(row.get("cost_price")) if row.get("cost_price") else None
        if row.get("my_price") and my_price is None:
            report.errors.append((line_no, _("Invalid my_price: %(v)s") % {"v": row.get("my_price")}))
            continue
        urls = [v for k, v in row.items() if k.startswith("url") and v]
        bad = []
        for url in urls:
            try:
                validate_url(url)
            except ValidationError:
                bad.append(url)
        if bad:
            report.errors.append((line_no, _("Invalid URL: %(v)s") % {"v": ", ".join(bad)}))
            continue
        try:
            with transaction.atomic():
                category = None
                if row.get("category"):
                    category, _c = Category.objects.get_or_create(organization=org, name=row["category"][:120])
                product, created = Product.objects.update_or_create(
                    organization=org,
                    sku=sku[:64],
                    defaults={"name": name[:255], "category": category, "my_price": my_price, "cost_price": cost_price},
                )
                new_urls = [u for u in urls if not product.offers.filter(url=u).exists()]
                if new_urls:
                    check_offer_limit(org, adding=len(new_urls))
                for url in new_urls:
                    CompetitorOffer.objects.create(
                        organization=org,
                        product=product,
                        store=_store_for(url),
                        url=url,
                        check_interval_minutes=default_interval,
                    )
                    report.created_offers += 1
                if created:
                    report.created_products += 1
                else:
                    report.updated_products += 1
        except LimitExceeded as exc:
            report.errors.append((line_no, str(exc)))
    return report
