import io
from decimal import Decimal

import pytest
from openpyxl import Workbook, load_workbook

from apps.accounts.models import Invitation, Membership, Role, User
from apps.alerts.models import AlertRule, TelegramLink
from apps.catalog.models import CompetitorOffer, Product, Store
from apps.pricing.models import PricingRule

from .factories import OfferFactory, ProductFactory, SnapshotFactory


@pytest.fixture
def product(org):
    p = ProductFactory(organization=org, my_price=Decimal(110_000), cost_price=Decimal(80_000))
    o = OfferFactory(product=p, last_price=Decimal(100_000), last_in_stock=True)
    SnapshotFactory(offer=o, price=Decimal(100_000))
    return p


PAGES = [
    "/",
    "/feed/",
    "/products/",
    "/products/new/",
    "/import/",
    "/alerts/",
    "/pricing/",
    "/billing/",
    "/scrapers/",
    "/telegram/",
    "/team/",
    "/settings/",
]


@pytest.mark.django_db
class TestPages:
    @pytest.mark.parametrize("url", PAGES)
    def test_pages_render(self, client_a, product, url):
        assert client_a.get(url).status_code == 200

    def test_product_pages(self, client_a, product):
        assert client_a.get(f"/products/{product.pk}/").status_code == 200
        assert client_a.get(f"/products/{product.pk}/edit/").status_code == 200
        data = client_a.get(f"/products/{product.pk}/chart.json?days=7").json()
        assert data["competitors"][0]["data"]
        assert data["mine"]["data"]

    def test_product_filters_and_htmx(self, client_a, product):
        resp = client_a.get("/products/?filter=expensive&q=Product", HTTP_HX_REQUEST="true")
        assert product.name in resp.content.decode()
        assert product.name not in client_a.get("/products/?filter=cheaper").content.decode()

    def test_anonymous_redirected(self, client, db):
        assert client.get("/").status_code == 302
        assert client.get("/login/").status_code == 200

    def test_user_without_org_redirected(self, client, db):
        user = User.objects.create_user("lonely@x.test", "pass12345!")
        client.force_login(user)
        assert client.get("/").url == "/org/new/"
        client.post("/org/new/", {"name": "New Co"})
        assert Membership.objects.filter(user=user, role=Role.OWNER).exists()

    def test_health_endpoint(self, client, db):
        assert client.get("/health/").json()["status"] == "ok"

    def test_language_switch(self, client_a):
        resp = client_a.post("/i18n/setlang/", {"language": "uz", "next": "/"})
        assert resp.status_code == 302


@pytest.mark.django_db
class TestActions:
    def test_signup_creates_org_and_trial(self, client):
        resp = client.post(
            "/signup/", {"email": "new@shop.uz", "password": "Very-Strong-123", "organization_name": "Shop UZ"}
        )
        assert resp.status_code == 302
        user = User.objects.get(email="new@shop.uz")
        org = user.memberships.get().organization
        assert org.subscription.plan.code == "free"
        assert AlertRule.objects.filter(organization=org).count() == 4

    def test_login(self, client, owner, org):
        resp = client.post("/login/", {"username": "owner@a.test", "password": "pass12345!"})
        assert resp.status_code == 302

    def test_create_product_and_offer(self, client_a, org, queued):
        resp = client_a.post(
            "/products/new/",
            {
                "sku": "NEW-1",
                "name": "New",
                "my_price": "1000",
                "cost_price": "500",
                "is_active": "on",
                "new_category": "TV",
            },
        )
        product = Product.objects.get(sku="NEW-1")
        assert resp.url == f"/products/{product.pk}/"
        assert product.category.name == "TV"
        client_a.post(
            f"/products/{product.pk}/offers/new/",
            {"url": "https://uzum.uz/ru/product/tv-123", "check_interval_minutes": 60},
        )
        offer = CompetitorOffer.objects.get(product=product)
        assert offer.store.domain == "uzum.uz"
        assert offer.store.adapter == Store.Adapter.UZUM
        assert queued.called

    def test_duplicate_sku_rejected(self, client_a, product):
        resp = client_a.post("/products/new/", {"sku": product.sku, "name": "dup", "my_price": "1"})
        assert resp.status_code == 200
        assert Product.objects.filter(sku=product.sku).count() == 1

    def test_edit_records_my_price_history(self, client_a, product):
        client_a.post(
            f"/products/{product.pk}/edit/",
            {"sku": product.sku, "name": product.name, "my_price": "95000", "cost_price": "80000", "is_active": "on"},
        )
        product.refresh_from_db()
        assert product.my_price == Decimal(95000)
        assert product.my_price_history.filter(price=Decimal(95000)).exists()

    def test_apply_recommended_price(self, client_a, product):
        client_a.post(f"/products/{product.pk}/apply-price/", {"price": "99000"})
        product.refresh_from_db()
        assert product.my_price == Decimal(99000)

    def test_check_now_htmx(self, client_a, product, queued):
        resp = client_a.post(f"/products/{product.pk}/check/", HTTP_HX_REQUEST="true")
        assert resp.status_code == 200 and queued.called

    def test_offer_toggle_and_delete(self, client_a, product):
        offer = product.offers.first()
        client_a.post(f"/offers/{offer.pk}/toggle/")
        offer.refresh_from_db()
        assert not offer.is_active
        client_a.post(f"/offers/{offer.pk}/toggle/")
        client_a.post(f"/offers/{offer.pk}/delete/")
        assert not CompetitorOffer.objects.filter(pk=offer.pk).exists()

    def test_delete_product(self, client_a, product):
        client_a.post(f"/products/{product.pk}/delete/")
        assert not Product.objects.filter(pk=product.pk).exists()

    def test_alert_rule_crud(self, client_a, org):
        client_a.post(
            "/alerts/",
            {"kind": "PRICE_CHANGED", "threshold": "0", "scope": "all", "notify_telegram": "on", "is_active": "on"},
        )
        rule = AlertRule.objects.get(organization=org, kind="PRICE_CHANGED")
        client_a.post(f"/alerts/{rule.pk}/toggle/")
        rule.refresh_from_db()
        assert not rule.is_active
        client_a.post(f"/alerts/{rule.pk}/delete/")
        assert not AlertRule.objects.filter(pk=rule.pk).exists()

    def test_alert_rule_scope_validation(self, client_a, org):
        resp = client_a.post("/alerts/", {"kind": "PRICE_CHANGED", "threshold": "0", "scope": "category"})
        assert resp.status_code == 200
        assert not AlertRule.objects.filter(organization=org, kind="PRICE_CHANGED").exists()

    def test_pricing_rule_crud(self, client_a, org):
        client_a.post(
            "/pricing/",
            {
                "name": "Avg",
                "strategy": "average",
                "amount": "0",
                "percent": "0",
                "min_margin_percent": "10",
                "round_to": "1000",
                "scope": "all",
                "priority": "1",
                "is_active": "on",
            },
        )
        rule = PricingRule.objects.get(organization=org, name="Avg")
        client_a.post(f"/pricing/{rule.pk}/delete/")
        assert not PricingRule.objects.filter(pk=rule.pk).exists()

    def test_viewer_is_read_only(self, client, viewer, product):
        client.force_login(viewer)
        assert client.get("/products/").status_code == 200
        assert client.post(f"/products/{product.pk}/delete/").status_code == 403
        assert client.get("/products/new/").status_code == 403
        assert client.post("/alerts/", {"kind": "PRICE_CHANGED", "scope": "all"}).status_code == 403

    def test_billing_request_creates_invoice(self, client_a, org):
        client_a.post("/billing/request/", {"plan": "start", "months": "3"})
        assert org.payments.get().months == 3

    def test_telegram_link_page(self, client_a, owner):
        client_a.get("/telegram/")
        link = TelegramLink.objects.get(user=owner)
        old = link.code
        client_a.post("/telegram/")
        link.refresh_from_db()
        assert link.code_expires_at is not None and (link.code != old or len(link.code) == 6)

    def test_settings_save(self, client_a, org, owner):
        client_a.post(
            "/settings/", {"name": "Renamed", "digest_enabled": "on", "digest_time": "08:30", "language": "uz"}
        )
        org.refresh_from_db()
        owner.refresh_from_db()
        assert org.name == "Renamed" and owner.language == "uz"

    def test_invite_flow(self, client_a, client, org):
        client_a.post("/team/", {"email": "colleague@a.test", "role": "manager"})
        invite = Invitation.objects.get(email="colleague@a.test")
        client.logout()
        assert client.get(f"/invite/{invite.token}/").status_code == 200
        resp = client.post(f"/invite/{invite.token}/", {"password": "Another-Pass-1"})
        assert resp.status_code == 302
        assert Membership.objects.filter(organization=org, user__email="colleague@a.test", role="manager").exists()

    def test_remove_member(self, client_a, viewer, org):
        m = Membership.objects.get(user=viewer)
        client_a.post(f"/team/{m.pk}/remove/")
        assert not Membership.objects.filter(pk=m.pk).exists()


@pytest.mark.django_db
class TestImportExport:
    def test_templates(self, client_a):
        assert client_a.get("/import/template.xlsx").status_code == 200
        assert b"sku,name" in client_a.get("/import/template.csv").content

    def test_csv_import_with_errors(self, client_a, org):
        csv_data = (
            b"sku;name;category;my_price;cost_price;url_1;url_2\n"
            b"A-1;Phone;Phones;1 299 000;900000;https://uzum.uz/ru/product/a-1;https://shop.example/p/1\n"
            b";No sku;;;;;\n"
            b"A-2;Bad url;;1000;;not-a-url;\n"
            b"A-3;Bad price;;abc;;;\n"
        )
        f = io.BytesIO(csv_data)
        f.name = "import.csv"
        resp = client_a.post("/import/", {"file": f, "check_interval_minutes": 60})
        report = resp.context["report"]
        assert report.created_products == 1 and report.created_offers == 2
        assert [line for line, _ in report.errors] == [3, 4, 5]
        p = Product.objects.get(organization=org, sku="A-1")
        assert p.my_price == Decimal(1299000)
        assert p.offers.get(url__contains="uzum").check_interval_minutes == 60

    def test_xlsx_import_updates_existing(self, client_a, org):
        ProductFactory(organization=org, sku="X-1", name="Old")
        wb = Workbook()
        ws = wb.active
        ws.append(["sku", "name", "my_price", "url_1"])
        ws.append(["X-1", "New name", 5000, "https://shop.example/x1"])
        buf = io.BytesIO()
        wb.save(buf)
        buf.seek(0)
        buf.name = "import.xlsx"
        report = client_a.post("/import/", {"file": buf, "check_interval_minutes": 360}).context["report"]
        assert report.updated_products == 1 and report.ok
        assert Product.objects.get(organization=org, sku="X-1").name == "New name"

    def test_import_respects_offer_limit(self, client_a, org):
        from apps.billing.models import Plan

        Plan.objects.filter(code="business").update(max_offers=1)
        org.subscription.refresh_from_db()
        f = io.BytesIO(b"sku,name,url_1,url_2\nL-1,Limited,https://a.example/1,https://b.example/2\n")
        f.name = "x.csv"
        report = client_a.post("/import/", {"file": f, "check_interval_minutes": 360}).context["report"]
        assert report.errors and "allows" in report.errors[0][1]
        assert not Product.objects.filter(sku="L-1").exists()

    def test_export_excel(self, client_a, product):
        resp = client_a.get("/export.xlsx")
        assert resp.status_code == 200
        wb = load_workbook(io.BytesIO(resp.content))
        ws = wb.active
        assert ws["A2"].value == product.sku
        assert ws.conditional_formatting  # "more expensive" highlight present

    def test_export_excel_requires_plan(self, client_a, org):
        from apps.billing.services import start_trial

        start_trial(org)
        resp = client_a.get("/export.xlsx")
        assert resp.status_code == 302 and resp.url == "/billing/"
        assert client_a.get("/export.csv").status_code == 200


@pytest.mark.django_db
def test_admin_store_test_view(client, product, respx_mock):
    import httpx

    admin = User.objects.create_superuser("root@x.test", "pass12345!")
    client.force_login(admin)
    store = product.offers.first().store
    assert client.get("/admin/catalog/store/").status_code == 200
    respx_mock.get(url__regex=r".*/robots\.txt").mock(return_value=httpx.Response(404))
    respx_mock.get("https://shop.test/p").mock(
        return_value=httpx.Response(
            200,
            text='<script type="application/ld+json">{"@type":"Product","name":"Z","offers":{"price":"777"}}</script>',
        )
    )
    resp = client.post(f"/admin/catalog/store/{store.pk}/test/", {"url": "https://shop.test/p"})
    assert resp.status_code == 200
    assert resp.context["result"].price == 777
